#!/usr/bin/env python
"""Serial device configuration and certificate service tool."""

import csv
import datetime
import os
import re
import sys
import tempfile
import time
import zipfile

import serial
import serial.tools.list_ports as ports
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PyQt5 import QtCore, QtWidgets, uic
from PyQt5.QtGui import QPixmap

BUILD_VERSION = '1.0.0'
EXCEL_FILENAME = 'cert_flashing_status.xlsx'
LEGACY_CSV_FILENAMES = (
    'cert_flashing_status.csv',
    'service_tool_device_checks.csv',
)
CERT_FILE_NAMES = ('ca.pem', 'cc.pem', 'ck.pem')


class FlowCancelled(Exception):
    pass


def get_excel_path():
    appDirectory = (
        os.path.dirname(os.path.abspath(sys.executable))
        if getattr(sys, 'frozen', False)
        else os.path.dirname(os.path.abspath(__file__))
    )
    return os.path.join(appDirectory, EXCEL_FILENAME)


class Worker(QtCore.QObject):
    logMsg = QtCore.pyqtSignal(str)
    errorMsg = QtCore.pyqtSignal(str)
    popupMsg = QtCore.pyqtSignal(str, str)
    progressChanged = QtCore.pyqtSignal(int)
    flowFinished = QtCore.pyqtSignal()

    def __init__(self):
        super(Worker, self).__init__()
        self.working = True
        self.serialPort = None
        self.scanRequested = False
        self.certificatePaths = []
        self.checksComplete = False
        self.results = {}
        self.expectedCertSizes = ()
        self.progressValue = 0

    def work(self):
        while self.working:
            try:
                if self.serialPort and self.serialPort.is_open:
                    if self.scanRequested:
                        self.scanRequested = False
                        self.runFlow()
                    else:
                        self.readNormalLog(0.05)
                else:
                    time.sleep(0.05)
            except (OSError, serial.SerialException) as error:
                message = 'Serial communication failed: {}'.format(error)
                self.logMsg.emit(message)
                self.popupMsg.emit('Serial communication error', message)
                if self.serialPort and self.serialPort.is_open:
                    self.serialPort.close()
                self.serialPort = None
                self.scanRequested = False
                self.checksComplete = False
                self.results = {}
                self.expectedCertSizes = ()
                self.progressChanged.emit(100)
                self.flowFinished.emit()
                time.sleep(0.1)
        if self.serialPort and self.serialPort.is_open:
            try:
                self.serialPort.close()
            except (OSError, serial.SerialException) as error:
                self.errorMsg.emit('Could not close the serial connection cleanly: {}'.format(error))
        self.serialPort = None

    def readLine(self, timeout):
        deadline = time.monotonic() + timeout
        while self.working and time.monotonic() < deadline:
            if self.serialPort.in_waiting:
                data = self.serialPort.readline().decode('utf-8', errors='replace')
                line = re.sub(r'\x1b\[\d+m', '', data)
                self.logMsg.emit(line.rstrip())
                return line
            time.sleep(0.01)
        return ''

    def readNormalLog(self, timeout):
        self.readLine(timeout)

    @staticmethod
    def isDefaultUin(uin):
        return bool(re.fullmatch(r'ACON4NA[X]+', uin.strip(), re.IGNORECASE))

    def runFlow(self):
        self.checksComplete = False
        self.results = {}
        self.progressValue = 0
        self.expectedCertSizes = ()
        self.progressChanged.emit(0)
        try:
            self.initializeResults()
            if not self.working:
                raise FlowCancelled()
            if self.scanStartupLog():
                self.logMsg.emit('Line automation device detected. Certificate flow skipped.')
                self.saveLineAutomationResult()
                self.popupMsg.emit(
                    'Default UIN Detected',
                    'UIN is default\n\nPut device in rejection bin',
                )
                return
            if not all(os.path.isfile(path) for path in self.certificatePaths):
                raise OSError('One or more certificate files are missing from resources.')
            self.flashCertificates()
            self.readCertificateStatus()
            self.collectDeviceInformation()
            if self.saveDeviceCheckResult():
                self.popupMsg.emit(
                    'Certificate Flashing',
                    'Certificate Flashing Successfully Completed',
                )
        except FlowCancelled:
            self.logMsg.emit('Device flow stopped because the application is closing.')
        except Exception as error:
            message = 'Device flow failed: {}'.format(error)
            self.logMsg.emit(message)
            self.popupMsg.emit('Device flow failed', str(error))
            self.saveDeviceCheckResult()
        finally:
            try:
                if self.serialPort and self.serialPort.is_open:
                    self.serialPort.close()
            except (OSError, serial.SerialException) as error:
                message = 'Could not close the serial connection cleanly: {}'.format(error)
                self.logMsg.emit(message)
                self.popupMsg.emit('Serial connection error', message)
            finally:
                self.resetCycleState()
                self.reportProgress(100)
                self.flowFinished.emit()

    def resetCycleState(self):
        self.serialPort = None
        self.scanRequested = False
        self.checksComplete = False
        self.results = {}
        self.expectedCertSizes = ()

    def reportProgress(self, value):
        self.progressValue = max(self.progressValue, min(100, value))
        self.progressChanged.emit(self.progressValue)

    def initializeResults(self):
        self.results = {
            'uin': '',
            'certs': '',
            'certificates_pass': False,
        }

    def scanStartupLog(self, timeout=3):
        self.logMsg.emit('Checking startup log for line automation...')
        deadline = time.monotonic() + timeout
        configPattern = re.compile(r'\$CONFIG\s*,\s*([^,\r\n]+)\s*,')
        while self.working and time.monotonic() < deadline:
            line = self.readLine(min(0.1, max(0, deadline - time.monotonic())))
            match = configPattern.search(line)
            if not match:
                continue
            uin = match.group(1).strip()
            self.results['uin'] = uin
            if self.isDefaultUin(uin):
                return True
            return False
        if not self.working:
            raise FlowCancelled()
        return False

    def collectDeviceInformation(self):
        self.logMsg.emit('Reading device information after certificate verification...')
        self.reportProgress(86)
        self.sendGetCommand(b'*GET#PRNCFG#')
        self.readPrncfg()
        self.reportProgress(99)
        self.checksComplete = True
        self.logMsg.emit('Device GET checks complete.')

    def waitWhileLogging(self, duration):
        deadline = time.monotonic() + duration
        while self.working and time.monotonic() < deadline:
            self.readLine(min(0.1, max(0, deadline - time.monotonic())))

    def sendGetCommand(self, command):
        if not self.working:
            raise FlowCancelled()
        self.logMsg.emit('TX: {}'.format(command.decode('ascii').strip()))
        self.serialPort.write(command)
        self.serialPort.flush()

    def readPrncfg(self):
        responseBuffer = ''
        deadline = time.monotonic() + 30
        while self.working and time.monotonic() < deadline:
            line = self.readLine(0.1)
            responseBuffer += line
            elapsed = 30 - max(0, deadline - time.monotonic())
            self.reportProgress(12 + int(elapsed / 30 * 8))
            match = re.search(r'\bUIN\s+([A-Za-z0-9_-]+)', responseBuffer, re.IGNORECASE)
            if match:
                self.results['uin'] = match.group(1)
                break
            if re.search(r'CALID\s+\S+\s*(?:\r?\n)+\s*=+', responseBuffer):
                break

        if not self.working:
            raise FlowCancelled()
        if not self.results['uin']:
            message = 'PRNCFG response is missing UIN.'
            self.logMsg.emit(message)
            self.popupMsg.emit('Device check incomplete', message)

    def readUntil(self, pattern, timeout, progressStart=None, progressEnd=None):
        buffer = ''
        deadline = time.monotonic() + timeout
        while self.working and time.monotonic() < deadline:
            buffer += self.readLine(0.1)
            if re.search(pattern, buffer):
                return buffer
            if progressStart is not None and progressEnd is not None:
                elapsed = timeout - max(0, deadline - time.monotonic())
                progress = progressStart + int(elapsed / timeout * (progressEnd - progressStart))
                self.reportProgress(progress)
        if not self.working:
            raise FlowCancelled()
        return buffer

    def flashCertificates(self):
        if len(self.certificatePaths) != 3:
            raise ValueError('Certificate files ca.pem, cc.pem, and ck.pem are required.')

        self.expectedCertSizes = tuple(os.path.getsize(path) for path in self.certificatePaths)
        totalSize = sum(self.expectedCertSizes)
        self.reportProgress(5)
        self.logMsg.emit('Starting certificates upload...')
        self.serialPort.reset_input_buffer()
        for certType, filePath, fileSize in zip(
            ('1', '2', '3'),
            self.certificatePaths,
            self.expectedCertSizes,
        ):
            certIndex = int(certType) - 1
            completedBefore = sum(self.expectedCertSizes[:certIndex])
            self.uploadCertificate(
                certType,
                filePath,
                fileSize,
                completedBefore,
                totalSize,
            )
            self.logMsg.emit('Finished uploading {}.'.format(os.path.basename(filePath)))
            if certType != '3':
                self.serialPort.flush()
                time.sleep(0.5)

        self.logMsg.emit('Certificate upload complete. Rebooting device...')
        self.serialPort.flush()
        time.sleep(0.5)
        self.reportProgress(65)
        self.serialPort.write(b'*SET#CRST#1#')
        self.logMsg.emit('Waiting for device to restart and resume normal logging...')
        self.waitForDeviceRestart(60)
        time.sleep(1)

    def readCertificateStatus(self):
        certStatusPattern = r'STATUS#CERT#\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*#'
        self.reportProgress(82)
        self.sendGetCommand(b'*GET#CERT#')
        response = self.readUntil(certStatusPattern, 30, 82, 86)
        match = re.search(certStatusPattern, response)
        if match:
            self.results['certs'] = ','.join(match.groups())
            actualCertSizes = tuple(int(size) for size in match.groups())
            self.results['certificates_pass'] = actualCertSizes == self.expectedCertSizes
            self.reportProgress(86)
            if not self.results['certificates_pass']:
                self.logMsg.emit(
                    'Certificate sizes mismatch: expected {}, received {}'.format(
                        self.expectedCertSizes,
                        actualCertSizes,
                    )
                )
                self.popupMsg.emit(
                    'Certificate Flashing Fail',
                    'Certificate Flashing Fail Connect device again',
                )
        else:
            responseDetails = response.strip()[-300:] or 'No response received.'
            message = 'Certificate status was not received after reboot. Response: {}'.format(
                responseDetails
            )
            self.logMsg.emit(message)
            self.popupMsg.emit(
                'Certificate Flashing Fail',
                'Certificate Flashing Fail Connect device again',
            )

    def uploadCertificate(self, certType, filePath, totalFileSize, completedBefore, allCertSize):
        if not self.working:
            raise FlowCancelled()
        initFrame = [
            '55AA', '3', '2', '0', '5.2.4', '1', certType, '0',
            totalFileSize, '0', '1024', 'FF',
        ]
        initPacket = ','.join(str(value) for value in initFrame).encode('ascii')
        self.logMsg.emit('Uploading {} ({} bytes)...'.format(os.path.basename(filePath), totalFileSize))
        time.sleep(0.1)
        self.serialPort.write(initPacket)
        self.serialPort.flush()

        receivedChunks = {}
        deadline = time.monotonic() + max(60, totalFileSize / 128)
        lastRequest = ''
        requestBuffer = ''
        with open(filePath, 'rb') as certFile:
            coveredBytes = 0
            while coveredBytes < totalFileSize:
                if not self.working:
                    raise FlowCancelled()
                if time.monotonic() >= deadline:
                    raise TimeoutError(
                        'Timed out while uploading {} (received {}/{} bytes; last device request: {}).'.format(
                            os.path.basename(filePath),
                            coveredBytes,
                            totalFileSize,
                            lastRequest.strip() or 'none',
                        )
                    )
                requestBuffer += self.readLine(0.2)
                requestMatch = re.search(
                    r'55AA,\s*\d+,\s*(\d+)\s*,\s*(\d+)\s*,',
                    requestBuffer,
                )
                if not requestMatch:
                    if '55AA' in requestBuffer:
                        lastRequest = requestBuffer[-300:]
                    if len(requestBuffer) > 1024:
                        requestBuffer = requestBuffer[-300:]
                    continue
                lastRequest = requestMatch.group(0)
                requestBuffer = requestBuffer[requestMatch.end():]
                offset, chunkSize = (int(value) for value in requestMatch.groups())
                if offset < 0 or offset >= totalFileSize or chunkSize <= 0:
                    raise ValueError('Invalid certificate chunk request: {}'.format(lastRequest.strip()))
                certFile.seek(offset)
                data = certFile.read(min(chunkSize, totalFileSize - offset))
                if not data:
                    raise ValueError('Device requested an empty certificate chunk.')

                crc = self.crc16(0, data)
                crcBytes = bytes([(crc >> 8) & 0xFF, crc & 0xFF])
                packet = (
                    b'55AA,4,'
                    + str(offset).encode('ascii')
                    + b','
                    + str(len(data)).encode('ascii')
                    + b','
                    + data
                    + b','
                    + crcBytes
                    + b',FF'
                )
                self.serialPort.write(packet)
                self.serialPort.flush()
                completed = completedBefore + min(totalFileSize, offset + len(data))
                self.reportProgress(10 + int(completed / allCertSize * 55))
                receivedChunks[offset] = max(
                    receivedChunks.get(offset, 0),
                    offset + len(data),
                )
                coveredBytes = 0
                for chunkOffset, chunkEnd in sorted(receivedChunks.items()):
                    if chunkOffset > coveredBytes:
                        break
                    coveredBytes = max(coveredBytes, chunkEnd)

    def waitForDeviceRestart(self, timeout):
        deadline = time.monotonic() + timeout
        while self.working and time.monotonic() < deadline:
            line = self.readLine(0.1)
            if re.search(r'UPTIME SEC 0,', line):
                self.reportProgress(82)
                return
            elapsed = timeout - max(0, deadline - time.monotonic())
            self.reportProgress(65 + int(elapsed / timeout * 17))
        if not self.working:
            raise FlowCancelled()
        raise TimeoutError('Timed out waiting for the device reboot.')

    @staticmethod
    def crc16(inputCrc, data):
        crc = inputCrc
        for value in data:
            crc ^= value << 8
            for _ in range(8):
                if crc & 0x8000:
                    crc = (crc << 1) ^ 0x1021
                else:
                    crc <<= 1
        return crc & 0xFFFF

    def saveDeviceCheckResult(self):
        uin = self.results.get('uin', '')
        uinPass = bool(uin)
        overallPass = (
            uinPass
            and self.results.get('certificates_pass', False)
        )
        failedChecks = []
        if not uinPass:
            failedChecks.append('UIN is missing or default')
        if not self.results.get('certificates_pass', False):
            failedChecks.append('certificate size verification did not pass')

        fieldnames = ['timestamp', 'uin', 'certs', 'result']
        row = {
            'timestamp': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'uin': uin,
            'certs': self.results.get('certs', ''),
            'result': 'PASS' if overallPass else 'FAIL',
        }
        if not self.saveExcelResult(fieldnames, row):
            return False

        self.logMsg.emit('{}: device checks saved to {}'.format(row['result'], get_excel_path()))
        if failedChecks:
            self.logMsg.emit('Failed checks: {}'.format('; '.join(failedChecks)))
        self.logMsg.emit(
            'uin: {}, certs: {}'.format(
                row['uin'],
                row['certs'] or 'not received',
            )
        )
        return overallPass

    def saveLineAutomationResult(self):
        fieldnames = ['timestamp', 'uin', 'certs', 'result']
        row = {
            'timestamp': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'uin': self.results.get('uin', ''),
            'certs': '',
            'result': 'FAIL',
        }
        return self.saveExcelResult(fieldnames, row)

    def saveExcelResult(self, fieldnames, row):
        try:
            excelPath = get_excel_path()
            self.appendExcelResult(excelPath, fieldnames, row)
        except (
            OSError,
            UnicodeError,
            csv.Error,
            ValueError,
            zipfile.BadZipFile,
            InvalidFileException,
        ) as error:
            message = 'Excel result write failed: {}'.format(error)
            self.logMsg.emit(message)
            self.popupMsg.emit('Excel write failed', message)
            return False
        self.logMsg.emit('Device result saved to {}'.format(excelPath))
        return True

    def appendExcelResult(self, excelPath, fieldnames, row):
        legacyRows = []
        if os.path.exists(excelPath):
            with open(excelPath, 'rb') as resultFile:
                isXlsx = resultFile.read(2) == b'PK'
            if isXlsx:
                workbook = load_workbook(excelPath)
            else:
                legacyRows = self.readLegacyCsv(excelPath, fieldnames)
                workbook = Workbook()
        else:
            workbook = Workbook()
            for legacyCsvFilename in LEGACY_CSV_FILENAMES:
                legacyCsvPath = os.path.join(
                    os.path.dirname(excelPath),
                    legacyCsvFilename,
                )
                if os.path.exists(legacyCsvPath):
                    legacyRows = self.readLegacyCsv(legacyCsvPath, fieldnames)
                    break

        tempPath = None
        try:
            worksheet = workbook.active
            existingRows = []
            if worksheet.max_row > 1 or worksheet.cell(row=1, column=1).value is not None:
                existingFields = [cell.value for cell in worksheet[1]]
                while existingFields and existingFields[-1] is None:
                    existingFields.pop()
                if (
                    not existingFields
                    or any(not isinstance(field, str) for field in existingFields)
                    or 'timestamp' not in existingFields
                ):
                    raise ValueError('The Excel result worksheet has an invalid header row.')
                for values in worksheet.iter_rows(min_row=2, values_only=True):
                    existingRow = dict(zip(existingFields, values))
                    normalizedRow = {
                        field: existingRow.get(field) or '' for field in fieldnames
                    }
                    if str(existingRow.get('Line Automation', '')).strip().lower() == 'detected':
                        normalizedRow['result'] = 'FAIL'
                    existingRows.append(normalizedRow)
                worksheet.delete_rows(1, worksheet.max_row)
            else:
                worksheet.delete_rows(1, worksheet.max_row)

            worksheet.append(fieldnames)
            for existingRow in existingRows + legacyRows:
                worksheet.append([existingRow.get(field, '') for field in fieldnames])
            worksheet.append([row.get(field, '') for field in fieldnames])

            fileDescriptor, tempPath = tempfile.mkstemp(
                suffix='.xlsx',
                dir=os.path.dirname(excelPath),
            )
            os.close(fileDescriptor)
            workbook.save(tempPath)
            os.replace(tempPath, excelPath)
        finally:
            if tempPath and os.path.exists(tempPath):
                os.remove(tempPath)
            workbook.close()

    @staticmethod
    def readLegacyCsv(csvPath, fieldnames):
        with open(csvPath, 'r', newline='', encoding='utf-8-sig') as csvFile:
            reader = csv.DictReader(csvFile)
            if not reader.fieldnames or 'timestamp' not in reader.fieldnames:
                raise ValueError(
                    'Cannot migrate {}: expected a legacy CSV header.'.format(csvPath)
                )
            migratedRows = []
            for row in reader:
                migratedRow = {field: row.get(field, '') or '' for field in fieldnames}
                if str(row.get('Line Automation', '')).strip().lower() == 'detected':
                    migratedRow['result'] = 'FAIL'
                migratedRows.append(migratedRow)
            return migratedRows


class ServiceToolUi(QtWidgets.QMainWindow):
    def __init__(self):
        super(ServiceToolUi, self).__init__()
        uic.loadUi(self.get_resource_path('service_tool.ui'), self)
        self.comboBoxComPort = self.findChild(QtWidgets.QComboBox, 'comboBoxComPort')
        self.btnConnect = self.findChild(QtWidgets.QPushButton, 'btnConnect')
        self.progressBar = self.findChild(QtWidgets.QProgressBar, 'progressBarFlow')
        self.labelLogo = self.findChild(QtWidgets.QLabel, 'labelLogo')
        self.headerSeparator = self.findChild(QtWidgets.QFrame, 'headerSeparator')
        self.log = self.findChild(QtWidgets.QTextBrowser, 'textBrowserLog')

        self.worker = Worker()
        self.worker.certificatePaths = [
            self.get_resource_path(os.path.join('resources', fileName))
            for fileName in CERT_FILE_NAMES
        ]
        logoPath = self.get_resource_path(os.path.join('resources', 'accolade-logo.png'))
        if not os.path.isfile(logoPath):
            raise FileNotFoundError('Logo image not found: {}'.format(logoPath))
        logoPixmap = QPixmap(logoPath)
        if logoPixmap.isNull():
            raise OSError('Could not load logo image: {}'.format(logoPath))
        samplePixmap = logoPixmap.scaled(256, 256, QtCore.Qt.IgnoreAspectRatio)
        sampleImage = samplePixmap.toImage()
        left, top = sampleImage.width(), sampleImage.height()
        right, bottom = -1, -1
        for y in range(sampleImage.height()):
            for x in range(sampleImage.width()):
                color = sampleImage.pixelColor(x, y)
                if color.alpha() and min(color.red(), color.green(), color.blue()) < 245:
                    left, top = min(left, x), min(top, y)
                    right, bottom = max(right, x), max(bottom, y)
        if right < left or bottom < top:
            raise OSError('Logo image contains no visible content: {}'.format(logoPath))
        sourceLeft = left * logoPixmap.width() // sampleImage.width()
        sourceTop = top * logoPixmap.height() // sampleImage.height()
        sourceRight = (right + 1) * logoPixmap.width() // sampleImage.width()
        sourceBottom = (bottom + 1) * logoPixmap.height() // sampleImage.height()
        visibleLogo = logoPixmap.copy(
            sourceLeft,
            sourceTop,
            sourceRight - sourceLeft,
            sourceBottom - sourceTop,
        )
        self.labelLogo.setPixmap(
            visibleLogo.scaled(
                self.labelLogo.size(),
                QtCore.Qt.KeepAspectRatio,
                QtCore.Qt.SmoothTransformation,
            )
        )
        self.logoPixmap = visibleLogo

        self.comboBoxComPort.addItems(
            '{}: {}'.format(port.device, port.description)
            for port in ports.comports()
        )
        self.btnConnect.clicked.connect(self.connectUart)

        self.progressResetTimer = QtCore.QTimer(self)
        self.progressResetTimer.setSingleShot(True)
        self.progressResetTimer.timeout.connect(lambda: self.progressBar.setValue(0))

        self.thread = QtCore.QThread(self)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.work)
        self.worker.logMsg.connect(self.logAppend)
        self.worker.errorMsg.connect(self.showError)
        self.worker.popupMsg.connect(self.showPopup)
        self.worker.progressChanged.connect(self.progressBar.setValue)
        self.worker.flowFinished.connect(self.resetForNextCycle)
        self.thread.start()
        self.closing = False
        self.setMinimumSize(1024, 640)

    def resizeEvent(self, event):
        super(ServiceToolUi, self).resizeEvent(event)
        if not hasattr(self, 'comboBoxComPort') or not hasattr(self, 'logoPixmap'):
            return

        width = self.centralWidget().width()
        height = self.centralWidget().height()
        margin = max(16, int(width * 0.02))
        headerHeight = max(88, int(height * 0.15))
        rowHeight = max(34, int(height * 0.045))
        progressHeight = max(26, int(height * 0.035))
        buttonWidth = max(100, int(width * 0.09))
        gap = max(28, int(width * 0.025))

        self.labelLogo.setGeometry(
            margin,
            margin,
            min(int(width * 0.23), 340),
            headerHeight - 2 * margin,
        )
        self.labelVersion.setGeometry(
            self.labelLogo.geometry().right() + gap,
            margin,
            width - self.labelLogo.geometry().right() - gap - margin,
            headerHeight - 2 * margin,
        )
        titleFont = self.labelVersion.font()
        titleFont.setPointSize(max(16, min(23, int(height * 0.024))))
        self.labelVersion.setFont(titleFont)
        self.labelVersion.setAlignment(QtCore.Qt.AlignVCenter | QtCore.Qt.AlignLeft)
        self.headerSeparator.setGeometry(margin, headerHeight, width - 2 * margin, 2)
        self.labelLogo.setPixmap(
            self.logoPixmap.scaled(
                self.labelLogo.size(),
                QtCore.Qt.KeepAspectRatio,
                QtCore.Qt.SmoothTransformation,
            )
        )

        controlsY = headerHeight + 2 * margin
        self.btnConnect.setGeometry(
            width - margin - buttonWidth,
            controlsY,
            buttonWidth,
            rowHeight,
        )
        self.comboBoxComPort.setGeometry(
            margin,
            controlsY,
            width - 3 * margin - buttonWidth,
            rowHeight,
        )
        controlsFont = self.comboBoxComPort.font()
        controlsFont.setPointSize(10)
        self.comboBoxComPort.setFont(controlsFont)
        self.btnConnect.setFont(controlsFont)
        progressY = controlsY + rowHeight + margin
        self.progressBar.setGeometry(
            margin,
            progressY,
            width - 2 * margin,
            progressHeight,
        )
        logY = progressY + progressHeight + margin
        self.log.setGeometry(
            margin,
            logY,
            width - 2 * margin,
            max(100, height - logY - margin),
        )
        progressFont = self.progressBar.font()
        progressFont.setPointSize(9)
        self.progressBar.setFont(progressFont)
        logFont = self.log.font()
        logFont.setPointSize(9)
        self.log.setFont(logFont)

    def connectUart(self):
        selectedPort = self.comboBoxComPort.currentText()
        portName = selectedPort.split(':', 1)[0].strip()
        if not portName or portName == 'Select Serial Port...':
            self.showError('Select a serial port before connecting.')
            return
        try:
            self.progressResetTimer.stop()
            self.progressBar.setValue(0)
            serialPort = serial.Serial(portName, 115200, timeout=0.1)
            self.worker.checksComplete = False
            self.worker.results = {}
            self.worker.scanRequested = True
            self.worker.serialPort = serialPort
        except (OSError, serial.SerialException) as error:
            self.worker.serialPort = None
            self.showError('Could not open {}: {}'.format(portName, error))
            return
        self.comboBoxComPort.setEnabled(False)
        self.btnConnect.setEnabled(False)
        self.btnConnect.setText('Connected')
        self.logAppend('Connected to {}. Starting device check and certificate flow...'.format(portName))

    def resetForNextCycle(self):
        self.log.clear()
        self.comboBoxComPort.setEnabled(True)
        self.btnConnect.setText('Connect')
        self.btnConnect.setEnabled(True)
        self.progressBar.setValue(100)
        self.progressResetTimer.start(1800)

    def get_resource_path(self, relativePath):
        if getattr(sys, 'frozen', False):
            return os.path.join(sys._MEIPASS, relativePath)
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), relativePath)

    def logAppend(self, value):
        self.log.append(value)

    def showError(self, message):
        QtWidgets.QMessageBox.critical(self, 'Service Tool', message)

    def showPopup(self, title, message):
        QtWidgets.QMessageBox.warning(self, title, message)

    def closeEvent(self, event):
        if self.thread.isRunning():
            if not self.closing:
                self.closing = True
                self.setEnabled(False)
                self.worker.working = False
                self.thread.finished.connect(self.close)
            self.thread.quit()
            event.ignore()
            return
        event.accept()


if __name__ == '__main__':
    app = QtWidgets.QApplication(sys.argv)
    window = ServiceToolUi()
    window.showMaximized()
    app.exec_()
